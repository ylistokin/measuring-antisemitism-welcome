"""Reproduce supplied Figure 2 aggregates; does not re-adjudicate articles."""
from pathlib import Path
import csv,json
from collections import Counter
import openpyxl
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

ROOT=Path(__file__).resolve().parents[1]
def main():
    workbook=openpyxl.load_workbook(ROOT/'data/racism_antisemitism_census_with_modern_figure_2.xlsx',data_only=True)
    sheet=workbook['Figure 2']
    rows=list(sheet.iter_rows(min_row=2,max_row=5,min_col=24,max_col=26,values_only=True))
    expected=[('Economics',27,2),('Political Science',13,4),('Psychology',46,2),('Sociology',77,4)]
    if rows!=expected:raise ValueError(f'Figure 2 source counts changed: {rows!r}')
    with (ROOT/'data/figure2_counts.csv').open('w',newline='') as f:
        writer=csv.writer(f);writer.writerow(['Discipline','Anti-Black','Antisemitism']);writer.writerows(rows)
    master=Counter()
    for row in workbook['Master Census'].iter_rows(min_row=6,values_only=True):
        if row[7]=='A' and row[8]=='Empirical':master[(row[0],row[6])]+=1
    for row in workbook['By Discipline'].iter_rows(min_row=6,values_only=True):
        if row[0] and master[(row[0],row[1])]!=row[4]:raise ValueError(f'Master/summary mismatch: {row[:5]}')
    black=sum(r[1] for r in rows);jewish=sum(r[2] for r in rows)
    summary={'figure2_totals':{'Anti-Black':black,'Antisemitism':jewish},'ratio':black/jewish,'unfiltered_empirical_A':{subject:sum(v for (d,s),v in master.items() if s==subject) for subject in ['Anti-Black','Antisemitism']},'modern_article_ledger_available':False}
    (ROOT/'data/figure2_summary.json').write_text(json.dumps(summary,indent=2)+'\n')
    fig,ax=plt.subplots(figsize=(12,7));xs=range(len(rows));width=.34
    ax.bar([x-width/2 for x in xs],[r[1] for r in rows],width,label='Anti-Black',color='#0072B2')
    ax.bar([x+width/2 for x in xs],[r[2] for r in rows],width,label='Antisemitism',color='#D55E00')
    ax.set_xticks(list(xs),[r[0] for r in rows]);ax.set_ylim(0,100);ax.set_ylabel('Total empirical papers')
    ax.set_axisbelow(True);ax.grid(axis='y',linestyle='--',alpha=.35);ax.legend(frameon=False,loc='upper center',ncol=2)
    for side in ['top','right']:ax.spines[side].set_visible(False)
    fig.suptitle('Figure 2: Modern Empirical Papers in Top-10 Journals, by Field',fontsize=16,fontweight='bold')
    fig.text(.08,.025,'Modern: at least one empirical observation after 1960; studies may begin earlier. Primary-focus empirical papers.\nPublication coverage: January 1, 2000–September 18, 2026; 2026 is partial. Source: supplied workbook.',fontsize=9)
    fig.tight_layout(rect=(0,.1,1,.94))
    for ext in ['png','svg','pdf']:fig.savefig(ROOT/f'figures/figure2-reproduced.{ext}',dpi=240)
    plt.close(fig);print(json.dumps(summary,indent=2))
if __name__=='__main__':main()
